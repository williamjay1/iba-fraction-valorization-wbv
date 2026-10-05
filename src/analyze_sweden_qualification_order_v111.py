"""Public-output reanalysis: qualification order in two Swedish road cases.

No process performance, acceptance fraction or statistical distribution is
imputed. Source model emissions are not independently rerun. All changed
reject, transport, water and disposal exchanges belong in signed Delta.
"""
from __future__ import annotations
import csv
import hashlib
import json
import math
import os
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(os.environ.get("PROJECT30_WORK", Path(__file__).resolve().parent.parent))
RAW = Path(os.environ.get("PROJECT30_SWEDISH_DOCX", ROOT / "original_files_not_distributed" / "swedish.docx"))
OUT = ROOT / "results" / "sweden_qualification_order_v111"

def save(name, rows):
    with (OUT/name).open("w", encoding="utf-8", newline="") as s:
        w=csv.DictWriter(s, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)

def main():
    if RAW.exists():
        from docx import Document
        input_path=RAW;digest=hashlib.sha256(RAW.read_bytes()).hexdigest()
        assert digest=='d76a4d4e4cd81f1024897abfacc04ebef55589679838de99cf759c3c4756e4e6'
        d=Document(RAW);assert len(d.tables)==5
        geometry=d.tables[1];results=d.tables[4]
        evidence_mode='original public supplementary DOCX'
    else:
        input_path=ROOT/'datasets'/'published_output_cells'/'swedish_source_table_cells.json'
        digest=hashlib.sha256(input_path.read_bytes()).hexdigest()
        meta=json.loads((input_path.parent/'source_cache_manifest.json').read_text(encoding='utf-8'))
        assert digest==meta['swedish_cell_export_sha256']
        assert meta['swedish_original_sha256']=='d76a4d4e4cd81f1024897abfacc04ebef55589679838de99cf759c3c4756e4e6'
        rows=json.loads(input_path.read_text(encoding='utf-8'))
        class TableCells:
            def __init__(self,rows):self.rows=rows
            def cell(self,r,c):return SimpleNamespace(text=self.rows[r][c])
        geometry=TableCells(rows['table_B1']);results=TableCells(rows['table_C1'])
        evidence_mode='derived public-table cell export, not original DOCX or full source LCA'
    assert "aggregate, case" in results.cell(3,0).text
    assert "low-intensity" in results.cell(8,0).text
    assert "high-intensity" in results.cell(9,0).text
    assert "HVO 100" in results.cell(17,0).text
    masses=[float(geometry.cell(r,1).text) for r in (15,18,21)]
    widths=[float(geometry.cell(r,1).text) for r in (1,2,3)]
    height=float(geometry.cell(5,1).text); density=float(geometry.cell(7,1).text)
    anchors=[]; threshold=[]; grid=[]; checks=[]; geometry_checks=[]
    for typ, mass, width in zip(("single","double","bike"),masses,widths):
        geometric=1000*width*height*density
        geometry_checks.append({"road_type":typ, "reported_MIBA_Mg_per_km":mass,
            "geometry_times_printed_density_Mg_per_km":geometric,
            "relative_mass_discrepancy_percent":100*(mass/geometric-1),
            "implied_density_Mg_per_m3":mass/(1000*width*height),
            "analysis_rule":"retain printed FU masses; flag discrepancy; do not silently replace"})
    # The source C1 reports one-at-a-time variations, not every factorial cell.
    for case in (1,2):
        for j,(typ,mass) in enumerate(zip(("single","double","bike"),masses)):
            col=(case-1)*3+j+1
            g_base=float(results.cell(3,col).text)*1000/mass
            for row,p in ((8,1.0),(9,10.0)):
                g_p=float(results.cell(row,col).text)*1000/mass
                err=g_p-g_base-p
                # Each printed result rounded independently to 0.1 t/km.
                maxerr=100/mass+1e-10
                assert abs(err)<=maxerr, (case,typ,p,err,maxerr)
                checks.append({"case":case,"road_type":typ,"processing_kgCO2eq_per_Mg":p,
                    "increment_from_source_kgCO2eq_per_Mg":g_p-g_base,
                    "error_kgCO2eq_per_Mg":err,"printed_rounding_bound":maxerr})
            for setting,row in (("HVO10_source_base",3),("HVO30",16),("HVO100",17)):
                g0=float(results.cell(row,col).text)*1000/mass
                assert g0<0
                # All-batch processing means p applies to every input Mg.
                # Eligible-only processing means p applies to q input Mg.
                benefit=-g0
                anchors.append({"case":case,"road_type":typ,"transport_setting":setting,
                    "source_FU":"1 km constructed road, source geometry",
                    "MIBA_Mg_per_km":mass,"source_net_tCO2eq_per_km":g0*mass/1000,
                    "benefit_kgCO2eq_per_Mg_MIBA":benefit,
                    "locator":f"Supplement Table B1 mass; C1 row {row}, column {col}"})
                for p in (0.0,1.0,10.0):
                    # Rounding intervals, not confidence intervals.
                    delta_b=50/mass
                    threshold.append({"case":case,"road_type":typ,"transport_setting":setting,
                        "processing_kgCO2eq_per_Mg":p,
                        "all_batch_q_break_even_if_Delta_zero":p/benefit,
                        "q_break_even_printed_rounding_lower":p/(benefit+delta_b),
                        "q_break_even_printed_rounding_upper":p/(benefit-delta_b),
                        "all_batch_feasible_for_q_at_most_one":p<=benefit,
                        "eligible_only_net_benefit_per_eligible_Mg":benefit-p,
                        "conditioning_order_penalty_at_q_half":0.5*p,
                        "interpretation":"q is a deterministic eligible diversion fraction, not measured regulatory acceptance; added burden scenario not an observed wash inventory"})
                    for k in range(101):
                        q=k/100
                        all_batch=q*g0+p
                        eligible=q*(g0+p)
                        assert math.isclose(all_batch-eligible,(1-q)*p,abs_tol=1e-10)
                        grid.append({"case":case,"road_type":typ,"transport_setting":setting,
                            "p":p,"q":q,"all_batch_net_if_Delta_zero":all_batch,
                            "eligible_only_net_if_Delta_zero":eligible,
                            "all_batch_Delta_max":-all_batch,"eligible_only_Delta_max":-eligible})
    OUT.mkdir(parents=True, exist_ok=True)
    for name,rows in (("source_anchors.csv",anchors),("qualification_order_thresholds.csv",threshold),
                      ("qualification_order_grid.csv",grid),("source_processing_rounding_checks.csv",checks),
                      ("source_geometry_check.csv",geometry_checks)):
        save(name,rows)
    m={"input":(input_path.relative_to(ROOT).as_posix() if input_path.is_relative_to(ROOT) else input_path.name),"input_sha256":digest,"source_doi":"10.1016/j.wasman.2026.115549",
       "source_evidence_mode":evidence_mode,
       "source_not_modified":hashlib.sha256(input_path.read_bytes()).hexdigest()==digest,
       "equations":{"all_batch":"G = q G_source + p + Delta", "eligible_only":"G = q (G_source + p) + Delta"},
       "Delta":"signed change from the assumed proportional source comparison; changed reject management, extra transport, water/sludge/residual handling and service-equivalence corrections",
       "assumptions":["linear proportional replacement of unchanged source road service",
                      "unqualified material follows unchanged source BAU unless accounted in Delta",
                      "conditioning p remains a scenario factor, not a measured process inventory"],
       "validation":"published processing increments reproduced within displayed rounding; NOT an independent rerun of the source LCA or field validation",
       "uncertainty":"epistemic q, p and Delta frontiers; source-displayed rounding intervals; no probability assigned to legal acceptance",
       "geometry_discrepancy_flagged":True}
    (OUT/"manifest.json").write_text(json.dumps(m,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"single_road_thresholds":[r for r in threshold if r['road_type']=='single'],
                      "geometry_checks":geometry_checks,"max_processing_reconstruction_error":max(abs(r['error_kgCO2eq_per_Mg']) for r in checks)},indent=2))

if __name__=="__main__":main()
